#!/usr/bin/env python3
"""The copy workbook: a spreadsheet CMS for every piece of site language.

    python3 copy_sheet.py export [copy/ericlombardi-ca-copy.xlsx]   # content files -> workbook
    python3 copy_sheet.py import [copy/ericlombardi-ca-copy.xlsx]   # "New copy" column -> content files

One tab per page. Columns: Key | Where it appears | Published copy | New copy | Notes.
Edit only "New copy". On import, any row whose New copy differs from Published copy is written
back to its source file (strings.json, site.json, consulting.json, support.json, or a Markdown
page), and the workbook is regenerated so Published copy matches the site again.
Blog posts are not in the workbook; they are edited as Markdown files in content/posts/.
"""
import json, pathlib, re, sys, datetime as dt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
C = HERE / "content"
DEFAULT = HERE / "copy" / "ericlombardi-ca-copy.xlsx"

TABS = ["Read me", "Site-wide", "Home", "Consulting", "Writing and media", "Article pages", "In the media", "Public policy",
        "Leadership 2026", "More Neighbours Toronto", "Build Toronto", "Events", "Support", "About", "Contact", "Privacy", "Not found"]
GROUP_TAB = {"site": "Site-wide", "home": "Home", "consulting": "Consulting", "writing": "Writing and media", "article": "Article pages",
             "media": "In the media", "policy": "Public policy", "events": "Events", "support": "Support", "about": "About",
             "contact": "Contact", "privacy": "Privacy", "notfound": "Not found"}
MD_PAGES = {"policy/_index.md": "Public policy", "policy/ontario-liberal-leadership-2026.md": "Leadership 2026",
            "policy/more-neighbours-toronto.md": "More Neighbours Toronto", "policy/build-toronto.md": "Build Toronto",
            "pages/about.md": "About", "pages/privacy.md": "Privacy"}
MD_FRONT = {"policy/_index.md": ["kicker", "headline", "lede"],
            "policy/ontario-liberal-leadership-2026.md": ["title", "card_title", "role", "dates", "summary", "feature_kicker", "feature_title", "feature_text", "feature_cta"],
            "policy/more-neighbours-toronto.md": ["title", "card_title", "role", "dates", "summary"],
            "policy/build-toronto.md": ["title", "card_title", "role", "dates", "summary"],
            "pages/about.md": ["headline", "lede"], "pages/privacy.md": ["headline", "lede"]}
FRONT_WHERE = {"title": "Page title", "card_title": "Title on the home page and public-policy record", "role": "Role (facts row and record)",
               "dates": "Dates (facts row and record)", "summary": "Summary beside the title, and on the record", "kicker": "Small label above the page title",
               "headline": "Page title", "lede": "Introduction beside the page title", "feature_kicker": "Platform feature, small label",
               "feature_title": "Platform feature heading", "feature_text": "Platform feature text", "feature_cta": "Platform feature button"}


# ----------------------------------------------------------------------------- content access
def jload(name): return json.loads((C / name).read_text(encoding="utf-8"))
def jsave(name, data): (C / name).write_text(json.dumps(data, ensure_ascii=False, indent=2 if name != "media.json" else 1) + "\n", encoding="utf-8")


def md_split(path):
    raw = (C / path).read_text(encoding="utf-8")
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n?(.*)$", raw, re.S)
    front = m.group(1) if m else ""
    body = m.group(2) if m else raw
    blocks = [b for b in re.split(r"\n\s*\n", body.strip("\n")) if b.strip()]
    parsed = []
    for b in blocks:
        lines = b.split("\n")
        if all(re.match(r"\s*[-*]\s+", l) for l in lines):
            parsed.append({"type": "bullets", "items": [re.sub(r"^\s*[-*]\s+", "", l) for l in lines]})
        elif lines[0].startswith("### "):
            parsed.append({"type": "h3", "items": [b[4:].strip()]})
        elif lines[0].startswith("## "):
            parsed.append({"type": "h2", "items": [b[3:].strip()]})
        elif lines[0].startswith("> "):
            parsed.append({"type": "quote", "items": [" ".join(l.lstrip("> ").strip() for l in lines)]})
        else:
            parsed.append({"type": "p", "items": [" ".join(l.strip() for l in lines)]})
    return front, parsed


def md_join(front, parsed):
    out = []
    for b in parsed:
        if b["type"] == "bullets": out.append("\n".join("- " + it for it in b["items"]))
        elif b["type"] == "h2": out.append("## " + b["items"][0])
        elif b["type"] == "h3": out.append("### " + b["items"][0])
        elif b["type"] == "quote": out.append("> " + b["items"][0])
        else: out.append(b["items"][0])
    return "---\n" + front + "\n---\n\n" + "\n\n".join(out) + "\n"


def front_get(front, key):
    m = re.search(rf"^{re.escape(key)}:\s*(.*)$", front, re.M)
    return m.group(1).strip() if m else None


def front_set(front, key, value):
    return re.sub(rf"^({re.escape(key)}:\s*).*$", lambda m: m.group(1) + value, front, count=1, flags=re.M)


TYPE_WHERE = {"p": "Paragraph", "h2": "Section heading", "h3": "Sub-heading", "bullets": "Bullet", "quote": "Quotation"}


# ----------------------------------------------------------------------------- rows
def rows():
    """Yield (tab, key, where, text) for every editable string."""
    strings = jload("strings.json")
    for g, items in strings.items():
        for k, v in items.items():
            yield GROUP_TAB[g], f"strings:{g}.{k}", v["where"], v["text"]
    s = jload("site.json")
    yield "Site-wide", "site:description", "Search-result description for the home page, the media kit short bio, and the Person record for search engines", s["description"]
    yield "Site-wide", "site:short_description", "Footer, under the name", s["short_description"]
    yield "Site-wide", "site:job_title", "Job title in the Person record for search engines", s["job_title"]
    for i, n in enumerate(s["nav"]):
        yield "Site-wide", f"site:nav[{i}].label", f"Navigation item {i + 1} (links to {n['href']})", n["label"]
    yield "Site-wide", "site:signup.title", "Sign-up block heading (home, writing, article pages, support)", s["signup"]["title"]
    yield "Site-wide", "site:signup.text", "Sign-up block paragraph", s["signup"]["text"]
    yield "Site-wide", "site:signup.statement", "Sign-up consent sentence under the form", s["signup"]["statement"]
    yield "Site-wide", "site:contact.headline", "Dark contact band heading (home) and Contact page title", s["contact"]["headline"]
    yield "Site-wide", "site:contact.text", "Dark contact band paragraph and Contact page introduction", s["contact"]["text"]
    yield "Home", "site:hero.kicker", "Opening, small label", s["hero"]["kicker"]
    yield "Home", "site:hero.headline", "Opening statement (large)", s["hero"]["headline"]
    yield "Home", "site:hero.lede", "Opening introduction", s["hero"]["lede"]
    for i, r in enumerate(s["hero"].get("record", [])):
        yield "Home", f"site:hero.record[{i}]", f"Record line under the buttons, item {i + 1}", r
    for i, p_ in enumerate(s["pillars"]):
        yield "Home", f"site:pillars[{i}].title", f"Dark band, area {i + 1} heading", p_["title"]
        yield "Home", f"site:pillars[{i}].text", f"Dark band, area {i + 1} text", p_["text"]
        yield "Home", f"site:pillars[{i}].cta", f"Dark band, area {i + 1} link", p_["cta"]
    for k, v in s["categories"].items():
        yield "Writing and media", f"site:categories.{k}", f"Category label for posts on this site ({k})", v
    c = jload("consulting.json")
    yield "Consulting", "consulting:kicker", "Small label above the page title", c["kicker"]
    yield "Consulting", "consulting:headline", "Page title", c["headline"]
    yield "Consulting", "consulting:lede", "Introduction beside the page title", c["lede"]
    for i, p_ in enumerate(c.get("intro", [])):
        yield "Consulting", f"consulting:intro[{i}]", f"Paragraph under the header, {i + 1}", p_
    for i, sv in enumerate(c["services"]):
        yield "Consulting", f"consulting:services[{i}].title", f"Dark band, service {i + 1} heading", sv["title"]
        yield "Consulting", f"consulting:services[{i}].summary", f"Dark band, service {i + 1} text", sv["summary"]
    for i, t in enumerate(c["sectors"]):
        yield "Consulting", f"consulting:sectors[{i}]", f"Sector tag {i + 1}", t
    yield "Consulting", "consulting:how.headline", "How I work, heading", c["how"]["headline"]
    yield "Consulting", "consulting:how.text", "How I work, introduction", c["how"]["text"]
    for i, st in enumerate(c["how"]["steps"]):
        yield "Consulting", f"consulting:how.steps[{i}].title", f"How I work, step {i + 1} heading", st["title"]
        yield "Consulting", f"consulting:how.steps[{i}].text", f"How I work, step {i + 1} text", st["text"]
    yield "Consulting", "consulting:partner.headline", "Partners heading", c["partner"]["headline"]
    for i, t in enumerate(c["partner"]["text"]):
        yield "Consulting", f"consulting:partner.text[{i}]", f"Partners paragraph {i + 1}", t
    for i, t in enumerate(c["partner"]["services"]):
        yield "Consulting", f"consulting:partner.services[{i}]", f"Partnership tag {i + 1}", t
    yield "Consulting", "consulting:cta.headline", "Dark band at the bottom, heading", c["cta"]["headline"]
    yield "Consulting", "consulting:cta.text", "Dark band at the bottom, text", c["cta"]["text"]
    sp = jload("support.json")
    yield "Support", "support:kicker", "Small label above the page title", sp["kicker"]
    yield "Support", "support:headline", "Page title", sp["headline"]
    yield "Support", "support:lede", "Introduction beside the page title", sp["lede"]
    yield "Support", "support:pledge.title", "Pledge panel heading", sp["pledge"]["title"]
    for i, t in enumerate(sp["pledge"]["paragraphs"]):
        yield "Support", f"support:pledge.paragraphs[{i}]", f"Pledge paragraph {i + 1}", t
    yield "Support", "support:contribute.title", "Contribute panel heading", sp["contribute"]["title"]
    yield "Support", "support:contribute.text", "Contribute panel text", sp["contribute"]["text"]
    yield "Support", "support:contribute.note", "Contribute panel fine print", sp["contribute"]["note"]
    for i, x in enumerate(sp["uses"]):
        yield "Support", f"support:uses[{i}].title", f"Where it goes, column {i + 1} heading", x["title"]
        yield "Support", f"support:uses[{i}].text", f"Where it goes, column {i + 1} text", x["text"]
    for i, q in enumerate(sp["faq"]):
        yield "Support", f"support:faq[{i}].q", f"Question {i + 1}", q["q"]
        yield "Support", f"support:faq[{i}].a", f"Answer {i + 1}", q["a"]
    for path, tab in MD_PAGES.items():
        front, parsed = md_split(path)
        for k in MD_FRONT[path]:
            v = front_get(front, k)
            if v is not None:
                yield tab, f"md:{path}:{k}", FRONT_WHERE.get(k, k), v
        for bi, b in enumerate(parsed):
            for ii, it in enumerate(b["items"]):
                key = f"md:{path}:body:{bi}" + (f".{ii}" if b["type"] == "bullets" else "")
                where = f"Body, {TYPE_WHERE[b['type']].lower()} {bi + 1}" + (f", item {ii + 1}" if b["type"] == "bullets" else "")
                yield tab, key, where, it


# ----------------------------------------------------------------------------- setters
def path_set(obj, path, value):
    """Set obj at a dotted path with [index] parts, e.g. hero.record[1] or services[2].title."""
    parts = re.findall(r"[^.\[\]]+|\[\d+\]", path)
    cur = obj
    for i, part in enumerate(parts):
        last = i == len(parts) - 1
        key = int(part[1:-1]) if part.startswith("[") else part
        if last:
            cur[key] = value
        else:
            cur = cur[key]


def apply(changes):
    """changes: list of (key, new_text). Returns a summary list."""
    done = []
    by_file = {}
    for key, new in changes:
        kind, rest = key.split(":", 1)
        by_file.setdefault(kind if kind != "md" else "md:" + rest.split(":")[0], []).append((key, new))
    for group, items in by_file.items():
        if group == "strings":
            data = jload("strings.json")
            for key, new in items:
                g, k = key.split(":", 1)[1].split(".", 1)
                data[g][k]["text"] = new; done.append(key)
            jsave("strings.json", data)
        elif group in ("site", "consulting", "support"):
            data = jload(group + ".json")
            for key, new in items:
                path_set(data, key.split(":", 1)[1], new); done.append(key)
            jsave(group + ".json", data)
        elif group.startswith("md:"):
            path = group[3:]
            front, parsed = md_split(path)
            for key, new in items:
                field = key.split(":", 2)[2]
                if field.startswith("body:"):
                    idx = field[5:].split(".")
                    b = parsed[int(idx[0])]
                    b["items"][int(idx[1]) if len(idx) > 1 else 0] = new
                else:
                    front = front_set(front, field, new)
                done.append(key)
            (C / path).write_text(md_join(front, parsed), encoding="utf-8")
    return done


# ----------------------------------------------------------------------------- workbook
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="172A31")
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
LOCK_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="D9D9D9")


def export(dest):
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    for name in TABS:
        ws = wb.create_sheet(name)
        sheets[name] = ws
        if name == "Read me":
            continue
        ws.append(["Key", "Where it appears", "Published copy", "New copy", "Notes"])
        for cell in ws[1]:
            cell.font = Font(name=FONT, bold=True, color="FFFFFF"); cell.fill = HEAD_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 36
        ws.column_dimensions["C"].width = 60; ws.column_dimensions["D"].width = 60; ws.column_dimensions["E"].width = 30
        ws.freeze_panes = "C2"
    counts = {}
    for tab, key, where, text in rows():
        ws = sheets[tab]
        ws.append([key, where, text, text, ""])
        r = ws.max_row
        counts[tab] = counts.get(tab, 0) + 1
        for col in range(1, 6):
            cell = ws.cell(row=r, column=col)
            cell.font = Font(name=FONT, size=10 if col in (1, 2) else 11, color="595959" if col in (1, 2) else "000000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        ws.cell(row=r, column=1).fill = LOCK_FILL; ws.cell(row=r, column=3).fill = LOCK_FILL
        ws.cell(row=r, column=4).fill = INPUT_FILL
    rm = sheets["Read me"]
    rm.column_dimensions["A"].width = 110
    lines = [
        ("ericlombardi.ca copy workbook", True),
        (f"Generated {dt.date.today():%B %d, %Y} from the site's content files. One tab per page.", False),
        ("", False),
        ("How to use it", True),
        ("Edit the yellow 'New copy' column only. 'Published copy' is what the site says now; 'Key' and 'Where it appears' tell the build where each line lives. Leave them alone.", False),
        ("When you are done, send the file back (or say the word) and the edits are imported into the site's content files. Rows where New copy equals Published copy are ignored.", False),
        ("After an import the workbook is regenerated, so Published copy catches up with the site and you can go again.", False),
        ("", False),
        ("Things to know", True),
        ("Text in {braces} such as {n}, {link} or {source} is filled in by the build (a count, a link, an outlet). Keep the token; move it if you like.", False),
        ("Body text on the About, Privacy and public-policy pages is Markdown: **bold** for emphasis, [text](url) for a link. Headings and bullets are separate rows; keep each row to one paragraph.", False),
        ("Adding or removing rows is not supported here. To add a bullet, a service area or a question, ask for it and it will be added to the content files.", False),
        ("Blog posts, the media list, events and the sign-up consent wording's legal review are outside this workbook.", False),
        ("Canadian spelling throughout. Sentence case for headings. No em dashes.", False),
        ("", False),
        ("Tabs and row counts", True),
    ] + [(f"{t}: {counts.get(t, 0)} rows", False) for t in TABS if t != "Read me"]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = rm.cell(row=i, column=1, value=txt)
        c.font = Font(name=FONT, bold=bold, size=12 if bold else 11); c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(dest)
    return sum(counts.values())


def norm(s):
    return re.sub(r"\s+", " ", (s or "")).strip()


def do_import(src):
    wb = load_workbook(src)
    current = {key: text for _, key, _, text in rows()}
    changes, skipped = [], []
    for ws in wb.worksheets:
        if ws.title == "Read me":
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            key, _, published, new = (row + (None,) * 5)[:4]
            if not key or key not in current:
                continue
            new_t = norm(new); pub_t = norm(published)
            if new_t == pub_t or not new_t:
                continue
            if pub_t != norm(current[key]):
                skipped.append((key, "the site changed since this workbook was exported; re-export and reapply"))
                continue
            changes.append((key, new.strip()))
    done = apply(changes)
    return done, skipped


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "export"
    target = pathlib.Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else DEFAULT
    if cmd == "export":
        n = export(target)
        print(f"wrote {target.relative_to(HERE) if target.is_relative_to(HERE) else target} ({n} rows)")
    elif cmd == "import":
        done, skipped = do_import(target)
        print(f"applied {len(done)} change(s)")
        for k in done: print("  updated", k)
        for k, why in skipped: print("  skipped", k, "->", why)
        if done:
            export(target)
            print("workbook regenerated; now run: python3 build.py")
    else:
        sys.exit(__doc__)
